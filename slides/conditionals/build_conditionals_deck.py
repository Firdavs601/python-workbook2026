"""Build the "Making Decisions" PowerPoint deck for the conditionals topic.

Starts from everyday decisions, then walks the docs in ``conditionals/_docs``:
comparison_operators, boolean_operators, if_elif_else, string_utils,
ascii_table, match.

The drawing primitives are the same ones used by the intro deck, so both
decks look like one course.

Run:  uv run --with python-pptx python slides/conditionals/build_conditionals_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).with_name("conditionals.pptx")

# ---------------------------------------------------------------- palette --
NAVY = RGBColor(0x1E, 0x2A, 0x38)
BLUE = RGBColor(0x37, 0x76, 0xAB)
YELLOW = RGBColor(0xFF, 0xD4, 0x3B)
GRAY = RGBColor(0x5A, 0x66, 0x72)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_FG = RGBColor(0xE8, 0xEE, 0xF4)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)

SANS = "Trebuchet MS"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.7)  # left/right margin
BODY_TOP = Inches(1.75)
BODY_W = W - 2 * M

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

_section = {"n": 0, "name": ""}


# ------------------------------------------------------------- primitives --
def _box(
    slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None
):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def _tf(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    return tf


def _para(
    tf,
    text,
    size,
    color,
    bold=False,
    font=SANS,
    first=False,
    space_before=0,
    space_after=0,
    align=PP_ALIGN.LEFT,
    italic=False,
):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if text:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
        f.color.rgb = color
    return p


def _rich(p, parts, size, default_color, font=SANS):
    """parts: list of (text, color|None, bold, mono)."""
    for text, color, bold, mono in parts:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.name = MONO if mono else font
        f.color.rgb = color or default_color


def _markup(text):
    """`code` -> mono+blue, *bold* -> bold navy."""
    parts, buf, i = [], "", 0
    while i < len(text):
        ch = text[i]
        if ch in "`*":
            end = text.find(ch, i + 1)
            if end > i:
                if buf:
                    parts.append((buf, None, False, False))
                    buf = ""
                inner = text[i + 1 : end]
                if ch == "`":
                    parts.append((inner, BLUE, False, True))
                else:
                    parts.append((inner, NAVY, True, False))
                i = end + 1
                continue
        buf += ch
        i += 1
    if buf:
        parts.append((buf, None, False, False))
    return parts


def slide(title=None, kicker=None, footer=True):
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=WHITE)
    if title is not None:
        _box(s, M, Inches(0.62), Inches(0.09), Inches(0.72), fill=YELLOW)
        tf = _tf(s, M + Inches(0.26), Inches(0.5), BODY_W - Inches(0.26), Inches(1.0))
        if kicker:
            _para(tf, kicker.upper(), 12, BLUE, bold=True, first=True, space_after=2)
            _para(tf, title, 30, NAVY, bold=True)
        else:
            _para(tf, title, 30, NAVY, bold=True, first=True)
    if footer and _section["name"]:
        tf = _tf(s, M, H - Inches(0.52), BODY_W, Inches(0.3))
        _para(tf, _section["name"], 10, GRAY, first=True)
    return s


def bullets(s, items, x=None, y=None, w=None, size=18, gap=9):
    """items: str, or (str, sub_level)."""
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    tf = _tf(s, x, y, w, H - y - Inches(0.7))
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        if text == "":
            _para(tf, "", 6, GRAY, first=first)
            first = False
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(0 if p is tf.paragraphs[0] else gap)
        p.space_after = Pt(0)
        p.level = lvl
        p.alignment = PP_ALIGN.LEFT
        glyph = "▸  " if lvl == 0 else "        –  "
        r = p.add_run()
        r.text = glyph
        r.font.size = Pt(size - (2 if lvl else 0))
        r.font.name = SANS
        r.font.color.rgb = BLUE if lvl == 0 else GRAY
        r.font.bold = lvl == 0
        _rich(p, _markup(text), size - (2 if lvl else 0), NAVY if lvl == 0 else GRAY)
    return tf


def code(s, lines, x=None, y=None, w=None, size=15, caption=None, dark=True):
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    lh = Pt(size * 1.42).emu
    h = Emu(int(lh * len(lines)) + Inches(0.46).emu)
    _box(
        s,
        x,
        y,
        w,
        h,
        fill=NAVY if dark else LIGHT,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.045,
    )
    tf = _tf(s, x + Inches(0.22), y + Inches(0.2), w - Inches(0.4), h - Inches(0.3))
    for i, ln in enumerate(lines):
        col = CODE_FG if dark else NAVY
        if ln.strip().startswith("#"):
            col = RGBColor(0x8E, 0xA6, 0xBD) if dark else GRAY
        _para(
            tf, ln if ln else " ", size, col, font=MONO, first=(i == 0), space_after=0
        )
    if caption:
        ctf = _tf(s, x, y + h + Inches(0.06), w, Inches(0.3))
        _para(ctf, caption, 13, GRAY, first=True, italic=True)
    return y + h + (Inches(0.42) if caption else Inches(0.24))


def out(s, lines, x=None, y=None, w=None, size=15, label="Output"):
    x = M if x is None else x
    w = BODY_W if w is None else w
    lh = Pt(size * 1.42).emu
    h = Emu(int(lh * len(lines)) + Inches(0.62).emu)
    _box(
        s,
        x,
        y,
        w,
        h,
        fill=LIGHT,
        line=RGBColor(0xD5, 0xDE, 0xE6),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.045,
    )
    tf = _tf(s, x + Inches(0.22), y + Inches(0.16), w - Inches(0.4), h - Inches(0.26))
    _para(tf, label.upper(), 10, GRAY, bold=True, first=True, space_after=3)
    for ln in lines:
        _para(tf, ln if ln else " ", size, NAVY, font=MONO)
    return y + h + Inches(0.24)


def table(
    s,
    headers,
    rows,
    x=None,
    y=None,
    w=None,
    size=14,
    col_w=None,
    mono_cols=(),
    height=None,
):
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    h = height or Inches(0.42 + 0.36 * len(rows))
    shp = s.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h)
    tbl = shp.table
    tbl.first_row = True
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(w * cw / total))
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = ""
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.margin_left = c.margin_right = Inches(0.1)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(c.text_frame, htxt, size, WHITE, bold=True, first=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            mono = j in mono_cols
            _rich(
                p, _markup(val) if not mono else [(val, BLUE, False, True)], size, NAVY
            )
    return y + h + Inches(0.2)


def note(s, text, y, kind="tip"):
    colors = {
        "tip": (BLUE, "TIP"),
        "warn": (RED, "WATCH OUT"),
        "ok": (GREEN, "REMEMBER"),
    }
    col, label = colors[kind]
    h = Inches(0.78)
    _box(s, M, y, BODY_W, h, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    _box(s, M, y, Inches(0.07), h, fill=col)
    tf = _tf(
        s,
        M + Inches(0.26),
        y + Inches(0.1),
        BODY_W - Inches(0.5),
        h - Inches(0.2),
        anchor=MSO_ANCHOR.MIDDLE,
    )
    _para(tf, label, 10, col, bold=True, first=True, space_after=2)
    p = tf.add_paragraph()
    _rich(p, _markup(text), 15, NAVY)
    return y + h + Inches(0.2)


def section(name, subtitle, points):
    _section["n"] += 1
    _section["name"] = name
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.22), W, Inches(0.22), fill=YELLOW)
    tf = _tf(s, M, Inches(2.1), Inches(8.4), Inches(3.0))
    _para(tf, f"PART {_section['n']}", 16, YELLOW, bold=True, first=True, space_after=8)
    _para(tf, name, 44, WHITE, bold=True, space_after=10)
    _para(tf, subtitle, 19, RGBColor(0xA9, 0xBD, 0xD1))
    tf2 = _tf(s, Inches(9.0), Inches(2.3), Inches(3.6), Inches(3.0))
    _para(tf2, "IN THIS PART", 11, YELLOW, bold=True, first=True, space_after=10)
    for pt in points:
        _para(tf2, "•  " + pt, 15, RGBColor(0xD6, 0xE2, 0xEC), space_after=7)
    return s


def flow(s, start, question, yes, no, y=None):
    """A read -> decide -> two outcomes diagram, the shape of every `if`."""
    y = BODY_TOP if y is None else y
    cy = y + Inches(0.7)

    def _label(shape_box, text, size=14, color=NAVY, bold=False):
        tf = _tf(
            s,
            shape_box[0],
            shape_box[1],
            shape_box[2],
            shape_box[3],
            anchor=MSO_ANCHOR.MIDDLE,
        )
        _para(tf, text, size, color, bold=bold, first=True, align=PP_ALIGN.CENTER)

    b = (M, cy - Inches(0.4), Inches(2.5), Inches(0.8))
    _box(
        s,
        *b,
        fill=LIGHT,
        line=RGBColor(0xC8, 0xD4, 0xDE),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12,
    )
    _label(b, start)

    _box(
        s,
        M + Inches(2.6),
        cy - Inches(0.13),
        Inches(0.5),
        Inches(0.26),
        fill=BLUE,
        shape=MSO_SHAPE.RIGHT_ARROW,
    )

    d = (M + Inches(3.25), cy - Inches(0.72), Inches(3.1), Inches(1.44))
    _box(s, *d, fill=YELLOW, shape=MSO_SHAPE.DIAMOND)
    _label(d, question, size=14, bold=True)

    _box(
        s,
        M + Inches(6.5),
        cy - Inches(0.13),
        Inches(0.6),
        Inches(0.26),
        fill=GREEN,
        shape=MSO_SHAPE.RIGHT_ARROW,
    )
    tfy = _tf(s, M + Inches(6.4), cy - Inches(0.52), Inches(0.8), Inches(0.3))
    _para(tfy, "yes", 12, GREEN, bold=True, first=True, align=PP_ALIGN.CENTER)

    yb = (M + Inches(7.25), cy - Inches(0.4), Inches(4.0), Inches(0.8))
    _box(
        s,
        *yb,
        fill=LIGHT,
        line=RGBColor(0xC8, 0xD4, 0xDE),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12,
    )
    _label(yb, yes, color=GREEN, bold=True)

    _box(
        s,
        M + Inches(4.67),
        cy + Inches(0.78),
        Inches(0.26),
        Inches(0.55),
        fill=RED,
        shape=MSO_SHAPE.DOWN_ARROW,
    )
    tfn = _tf(s, M + Inches(5.0), cy + Inches(0.9), Inches(0.8), Inches(0.3))
    _para(tfn, "no", 12, RED, bold=True, first=True)

    nb = (M + Inches(2.8), cy + Inches(1.42), Inches(4.0), Inches(0.8))
    _box(
        s,
        *nb,
        fill=LIGHT,
        line=RGBColor(0xC8, 0xD4, 0xDE),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12,
    )
    _label(nb, no, color=RED, bold=True)

    return cy + Inches(2.42)


# ================================================================== TITLE ==
def title_slide():
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, 0, Inches(0.26), H, fill=YELLOW)
    tf = _tf(s, Inches(1.3), Inches(1.9), Inches(10.5), Inches(3.4))
    _para(
        tf,
        "CONDITIONALS MODULE  ·  PYTHON WORKBOOK",
        15,
        YELLOW,
        bold=True,
        first=True,
        space_after=14,
    )
    _para(tf, "Making Decisions", 54, WHITE, bold=True, space_after=4)
    _para(
        tf,
        "if, elif, else — and everything they need",
        30,
        RGBColor(0x8F, 0xB6, 0xD9),
        space_after=22,
    )
    _para(
        tf,
        "Your programs stop running in a straight line and start choosing what to do.",
        18,
        RGBColor(0xC3, 0xD3, 0xE1),
    )
    tf2 = _tf(s, Inches(1.3), Inches(6.0), Inches(10.5), Inches(0.9))
    _para(
        tf2,
        "comparison_operators  ·  boolean_operators  ·  "
        "if_elif_else  ·  string_utils  ·  ascii_table  "
        "·  match",
        13,
        RGBColor(0x7E, 0x93, 0xA6),
        font=MONO,
        first=True,
    )


def roadmap_slide():
    s = slide("What we will cover today", footer=False)
    rows = [
        (
            "1",
            "Decisions in Everyday Life",
            'From "if it rains, take an umbrella" to your first if statement',
        ),
        (
            "2",
            "Comparison Operators",
            "==  !=  <  <=  >  >=  and the True / False they produce",
        ),
        (
            "3",
            "Booleans: and, or, not",
            "Truth tables, precedence, short-circuit evaluation",
        ),
        (
            "4",
            "if / elif / else",
            "The workhorse of this chapter, plus nesting and order of checks",
        ),
        (
            "5",
            "Conditions on Text",
            "Comparing letters and words, ord / chr, the ASCII table",
        ),
        (
            "6",
            "The match Statement",
            "Clean branching when one value has many possible answers",
        ),
    ]
    y = BODY_TOP
    for num, name, desc in rows:
        _box(
            s,
            M,
            y,
            Inches(0.62),
            Inches(0.62),
            fill=BLUE,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.15,
        )
        tfn = _tf(s, M, y, Inches(0.62), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
        _para(tfn, num, 20, WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
        tft = _tf(s, M + Inches(0.85), y + Inches(0.02), Inches(10.5), Inches(0.62))
        _para(tft, name, 19, NAVY, bold=True, first=True, space_after=1)
        _para(tft, desc, 14, GRAY)
        y += Inches(0.75)
    note(
        s,
        "Everything here is written down in `conditionals/_docs/` — "
        "the slides are the guided tour, the docs are the reference.",
        y + Inches(0.05),
        "tip",
    )


# ============================================ PART 1: DECISIONS IN LIFE ==
def part1():
    section(
        "Decisions in Everyday Life",
        "You already know how conditionals work — today you learn to write them down",
        [
            "Straight-line programs",
            "Decisions you make daily",
            "One rule, three notations",
            "Your first if",
            "else",
            "Indentation",
        ],
    )

    s = slide("Where we are now: programs run in a straight line", kicker="The problem")
    y = code(
        s,
        [
            "radius = float(input())",
            "area = 3.14159 * radius ** 2",
            'print(f"{area:.2f}")',
        ],
        caption="Every line runs, always, top to bottom. Nothing is ever skipped.",
    )
    bullets(
        s,
        [
            "That is enough for a calculator, but not for a *decision*.",
            'This program cannot say "a negative radius makes no sense".',
            "It cannot answer differently for different input — it only computes.",
            "Today we let the program *choose* which lines to run.",
        ],
        y=y + Inches(0.1),
        size=18,
    )

    s = slide(
        "You make dozens of these decisions before breakfast", kicker="Everyday logic"
    )
    bullets(
        s,
        [
            "*If* it is raining, take an umbrella. *Otherwise*, take sunglasses.",
            "*If* the traffic light is green, cross the street. *Otherwise*, wait.",
            "*If* you are 18 or older, you may vote.",
            "*If* the bill is over 100 TJS, delivery is free.",
            "*If* the bus is full, wait for the next one.",
            "*If* your phone battery is under 20%, look for a charger.",
        ],
        size=20,
        gap=13,
    )
    note(
        s,
        "Each one has the same shape: a *question* with a yes/no answer, "
        "and a *different action* for each answer.",
        Inches(5.6),
        "ok",
    )

    s = slide("The shape of a decision", kicker="Everyday logic")
    y = flow(
        s, "Look outside", "Is it\nraining?", "Take an umbrella", "Take sunglasses"
    )
    note(
        s,
        "A diamond is a question that can only be answered *yes* or *no*. "
        "Programmers call that answer a *boolean*.",
        y + Inches(0.05),
        "tip",
    )

    s = slide("The same rule, written three ways", kicker="Everyday logic")
    tf = _tf(s, M, Inches(1.72), BODY_W, Inches(0.4))
    _rich(
        tf.paragraphs[0],
        _markup(
            "*In words:* if the person is 18 or older, they can vote; "
            "otherwise they cannot."
        ),
        18,
        NAVY,
    )
    y = flow(
        s,
        "Read age",
        "age >= 18 ?",
        "You can vote",
        "You cannot vote yet",
        y=Inches(2.15),
    )
    code(
        s,
        [
            "age = int(input())",
            "if age >= 18:",
            '    print("You can vote")',
            "else:",
            '    print("You cannot vote yet")',
        ],
        y=y + Inches(0.05),
        size=14,
    )

    s = slide("Your first if statement, line by line", kicker="Syntax")
    y = code(
        s,
        [
            "if age >= 18:",
            '    print("You can vote")',
        ],
        size=20,
    )
    table(
        s,
        ["Part", "What it is", "Rule"],
        [
            ("`if`", "The keyword that starts a decision", "Always lowercase"),
            (
                "`age >= 18`",
                "The condition — a yes/no question",
                "Must evaluate to `True` or `False`",
            ),
            ("`:`", "Colon at the end of the line", "Forgetting it is a `SyntaxError`"),
            (
                "`    print(...)`",
                "The body — what to do if the answer is yes",
                "Indented by 4 spaces",
            ),
        ],
        y=y + Inches(0.1),
        col_w=[3, 6, 5],
        size=15,
    )

    s = slide("Adding the other branch: else", kicker="Syntax")
    y = code(
        s,
        [
            "temperature = int(input())",
            "",
            "if temperature > 30:",
            '    print("Stay in the shade")',
            "else:",
            '    print("Enjoy the weather")',
        ],
    )
    y = out(s, ["35", "Stay in the shade"], y=y, label="Run 1")
    out(s, ["22", "Enjoy the weather"], y=y, label="Run 2")

    s = slide("Indentation is not decoration — it is the block", kicker="Syntax")
    y = code(
        s,
        [
            "if score > 90:",
            '    print("Excellent")      # inside the if: only for score > 90',
            '    print("Well done")      # also inside',
            'print("Goodbye")            # outside: printed every time',
        ],
    )
    y = note(
        s,
        "Other languages use braces. Python uses *spaces*, so wrong "
        "indentation silently changes what your program means.",
        y,
        "warn",
    )
    bullets(
        s,
        [
            "Use *4 spaces* per level. Let your editor insert them with Tab.",
            "Never mix tabs and spaces in one file — Python will refuse to run it.",
        ],
        y=y + Inches(0.05),
        size=17,
    )


# ======================================= PART 2: COMPARISON OPERATORS ==
def part2():
    section(
        "Comparison Operators",
        "The six questions you can ask about two values",
        [
            "The six operators",
            "Results are booleans",
            "= vs ==",
            "Comparing text",
            "Chained comparisons",
        ],
    )

    s = slide("The six comparison operators", kicker="Operators")
    table(
        s,
        ["Operator", "Name", "Returns `True` when"],
        [
            ("`==`", "Equal to", "Both values are equal"),
            ("`!=`", "Not equal to", "The values are different"),
            ("`<`", "Less than", "The left value is smaller"),
            ("`<=`", "Less than or equal to", "The left value is smaller or the same"),
            ("`>`", "Greater than", "The left value is bigger"),
            (
                "`>=`",
                "Greater than or equal to",
                "The left value is bigger or the same",
            ),
        ],
        col_w=[2, 4, 7],
        size=16,
    )
    note(
        s,
        "`<=` and `>=` are written in that order, and never with a space in between.",
        Inches(4.9),
        "warn",
    )

    s = slide("Every comparison is worth seeing for yourself", kicker="Operators")
    table(
        s,
        ["Operator", "Example", "Result"],
        [
            (
                "`==`",
                '`5 == 5`   `5 == 3`   `"hi" == "hi"`',
                "`True`   `False`   `True`",
            ),
            (
                "`!=`",
                '`5 != 3`   `5 != 5`   `"hi" != "ho"`',
                "`True`   `False`   `True`",
            ),
            ("`<`", "`3 < 5`   `5 < 3`   `5 < 5`", "`True`   `False`   `False`"),
            ("`<=`", "`3 <= 5`   `5 <= 5`   `7 <= 5`", "`True`   `True`   `False`"),
            ("`>`", "`5 > 3`   `3 > 5`   `5 > 5`", "`True`   `False`   `False`"),
            ("`>=`", "`5 >= 3`   `5 >= 5`   `3 >= 5`", "`True`   `True`   `False`"),
        ],
        col_w=[2, 7, 5],
        size=15,
    )
    note(
        s,
        "Type them into Python yourself. Two minutes at the interpreter "
        "beats ten minutes of staring at a table.",
        Inches(5.1),
        "tip",
    )

    s = slide("A condition is a value you can print", kicker="Booleans")
    y = code(
        s,
        [
            "print(5 > 3)            # True",
            "print(5 < 3)            # False",
            "print(type(5 > 3))      # <class 'bool'>",
            "",
            "is_adult = age >= 18    # store the answer in a variable",
            "print(is_adult)         # True or False",
        ],
    )
    bullets(
        s,
        [
            "A comparison does not print anything — it *produces a value*.",
            "That value has its own type: `bool`, with exactly two members.",
            "`if` simply looks at that value and decides whether to run the body.",
        ],
        y=y + Inches(0.1),
        size=18,
    )

    s = slide("The classic beginner bug: `=` vs `==`", kicker="Booleans")
    y = code(
        s,
        [
            "age = 18        # assignment: put 18 into the variable age",
            "age == 18       # question: is age equal to 18?  -> True",
            "",
            "if age = 18:    # SyntaxError: Python refuses this",
            "if age == 18:   # correct",
        ],
    )
    note(
        s,
        "Read `=` as *becomes* and `==` as *is equal to*. Say it out loud "
        "while you type — it works.",
        y,
        "ok",
    )

    s = slide("Comparing text", kicker="Operators")
    y = code(
        s,
        [
            'print("cat" == "cat")      # True',
            'print("Cat" == "cat")      # False  - capital C is a different character',
            'print("apple" < "banana")  # True   - alphabetical order',
            "",
            'month = input().lower()   # so "March", "MARCH" and "march" all match',
            'if month == "march":',
            '    print("Spring is close")',
        ],
        size=14,
    )
    note(
        s,
        "Users type what they like. Normalising with `.lower()` before "
        "comparing saves you from a dozen `elif` branches.",
        y,
        "tip",
    )

    s = slide("Chained comparisons: ranges the readable way", kicker="Operators")
    y = code(
        s,
        [
            "# Is the BMI in the normal range?",
            "if 18.5 <= bmi < 25.0:",
            '    print("Normal weight")',
            "",
            "# The same thing spelled out (Part 3 explains `and`)",
            "if bmi >= 18.5 and bmi < 25.0:",
            '    print("Normal weight")',
        ],
    )
    bullets(
        s,
        [
            "Python lets you write a range exactly as it looks in mathematics.",
            "`18.5 <= bmi < 25.0` includes 18.5 and excludes 25.0 — read the "
            "table in the task carefully.",
            "Most exercises in this chapter are built from ranges like this one.",
        ],
        y=y + Inches(0.1),
        size=17,
    )


# =========================================== PART 3: BOOLEAN OPERATORS ==
def part3():
    section(
        "Booleans: and, or, not",
        "Building bigger questions out of small ones",
        ["The bool type", "and", "or", "not", "Precedence", "Short-circuit"],
    )

    s = slide("The boolean type", kicker="bool")
    bullets(
        s,
        [
            "`bool` has only *two* values: `True` and `False`.",
            "Capital `T` and capital `F` — `true` is not a Python keyword.",
            "Every comparison produces one of them.",
            "`and`, `or` and `not` combine them into more complex conditions.",
        ],
        size=19,
    )
    y = code(
        s,
        [
            "raining = True",
            "cold = False",
            "print(raining and cold)   # False",
            "print(raining or cold)    # True",
            "print(not raining)        # False",
        ],
        y=Inches(3.9),
    )
    note(
        s,
        "These three words are the whole vocabulary. Everything else is combination.",
        y,
        "ok",
    )

    s = slide("and — true only when both are true", kicker="and")
    table(
        s,
        ["A", "B", "`A and B`"],
        [
            ("`True`", "`True`", "`True`"),
            ("`True`", "`False`", "`False`"),
            ("`False`", "`True`", "`False`"),
            ("`False`", "`False`", "`False`"),
        ],
        w=Inches(5.2),
        col_w=[1, 1, 2],
        size=16,
    )
    y = code(
        s,
        [
            "# A triangle is valid only if all three conditions hold",
            "if a + b > c and a + c > b and b + c > a:",
            '    print("valid triangle")',
            "",
            "# You may enter only with a ticket AND an ID",
            "if has_ticket and has_id:",
            '    print("Welcome")',
        ],
        x=M + Inches(5.6),
        w=BODY_W - Inches(5.6),
        size=13,
    )
    note(
        s,
        "Think of `and` as a checklist: one unticked box and the whole "
        "condition is `False`.",
        Inches(4.6),
        "tip",
    )

    s = slide("or — true when at least one is true", kicker="or")
    table(
        s,
        ["A", "B", "`A or B`"],
        [
            ("`True`", "`True`", "`True`"),
            ("`True`", "`False`", "`True`"),
            ("`False`", "`True`", "`True`"),
            ("`False`", "`False`", "`False`"),
        ],
        w=Inches(5.2),
        col_w=[1, 1, 2],
        size=16,
    )
    y = code(
        s,
        [
            "# Two of the three sides equal -> isosceles",
            "if a == b or a == c or b == c:",
            '    print("isosceles")',
            "",
            "# Weekend: Saturday or Sunday",
            'if day == "saturday" or day == "sunday":',
            '    print("No classes today")',
        ],
        x=M + Inches(5.6),
        w=BODY_W - Inches(5.6),
        size=13,
    )
    note(
        s,
        '`or` in Python is the *inclusive* or: "tea or coffee" here also allows both.',
        Inches(4.6),
        "warn",
    )

    s = slide("not — flip the answer", kicker="not")
    table(
        s,
        ["A", "`not A`"],
        [("`True`", "`False`"), ("`False`", "`True`")],
        w=Inches(4.2),
        col_w=[1, 1],
        size=16,
    )
    y = code(
        s,
        [
            "if not is_open:",
            '    print("Come back tomorrow")',
            "",
            "# Often the clearer form is the positive one:",
            "if is_closed:",
            '    print("Come back tomorrow")',
        ],
        x=M + Inches(4.6),
        w=BODY_W - Inches(4.6),
        size=14,
    )
    note(
        s,
        "`not` is useful, but a condition with two or three `not`s in it "
        "is a condition nobody can read — including you, next week.",
        Inches(4.3),
        "warn",
    )

    s = slide("Precedence: which operator goes first", kicker="Precedence")
    table(
        s,
        ["Order", "Operator", "Meaning"],
        [
            ("1 (highest)", "`not`", "Applied to the value right after it"),
            ("2", "`and`", "Binds tighter than `or`"),
            ("3 (lowest)", "`or`", "Applied last"),
        ],
        col_w=[3, 3, 8],
        size=16,
    )
    y = code(
        s,
        [
            "result = not True or False and True",
            "print(result)",
            "# Is evaluated as: (not True) or (False and True)",
            "# Result: False or False -> False",
        ],
        y=Inches(3.5),
    )
    note(
        s,
        "Do not memorise the table — add *brackets*. "
        "`(not raining) and (temperature > 20)` is never ambiguous.",
        y,
        "ok",
    )

    s = slide("Short-circuit evaluation", kicker="Short-circuit")
    bullets(
        s,
        [
            "Python reads a condition left to right and stops as soon as the "
            "answer is certain.",
            "`and`: stops if the left side is `False` — the rest cannot "
            "change the outcome.",
            "`or`: stops if the left side is `True` — it is already satisfied.",
        ],
        size=18,
    )
    y = code(
        s,
        [
            "x = 0",
            "if x != 0 and (10 / x) > 1:",
            '    print("Valid")',
            "# x != 0 is False, so Python stops and never attempts 10 / x",
            "# Without short-circuiting, this would crash with ZeroDivisionError",
        ],
        y=Inches(3.7),
        size=14,
    )
    note(
        s,
        "This is why the *guard* goes first: check that a value is safe "
        "before you use it.",
        y,
        "ok",
    )


# ================================================= PART 4: IF / ELIF / ELSE ==
def part4():
    section(
        "if / elif / else",
        "One question, two questions, many questions",
        [
            "The three shapes",
            "Worked examples",
            "Order of checks",
            "elif vs many ifs",
            "Nesting",
        ],
    )

    s = slide("Shape 1 and 2: if, and if / else", kicker="Structure")
    y = code(
        s,
        [
            "# 1. if        - do something, or nothing at all",
            "if temperature > 30:",
            '    print("Hot day")',
            "",
            "# 2. if / else - exactly one of two branches runs",
            "if age >= 18:",
            '    print("You can vote")',
            "else:",
            '    print("You cannot vote yet")',
        ],
    )
    y = note(
        s,
        "A lone `if` may do nothing at all. An `if` / `else` always "
        "does exactly one of the two things.",
        y,
        "ok",
    )
    bullets(
        s,
        [
            'Use a lone `if` when "nothing happens" is a valid outcome.',
            "Use `if` / `else` when the task always expects some output.",
        ],
        y=y + Inches(0.05),
        size=17,
    )

    s = slide("Shape 3: if / elif / else", kicker="Structure")
    y = code(
        s,
        [
            "# pick one branch out of many",
            "if num > 0:",
            '    print("positive")',
            "elif num < 0:",
            '    print("negative")',
            "else:",
            '    print("zero")',
        ],
    )
    y = note(
        s,
        "`elif` is short for *else if*. You may write as many `elif` "
        "branches as you need; `else` is optional and comes last.",
        y,
        "ok",
    )
    bullets(
        s,
        [
            "Python checks the branches *top to bottom* and takes the first "
            "`True` one.",
            "Everything below the winning branch is skipped, even if it is also true.",
            "`else` is the catch-all: it runs when nothing above it matched.",
        ],
        y=y + Inches(0.05),
        size=17,
    )

    s = slide("Example 1: can a person vote?", kicker="Worked example")
    y = code(
        s,
        [
            "# Read age from user",
            'age = int(input("Enter your age: "))',
            "",
            "# Check if eligible to vote",
            "if age >= 18:",
            '    print("You can vote")',
            "else:",
            '    print("You cannot vote yet")',
        ],
    )
    y = out(s, ["17", "You cannot vote yet"], y=y)
    note(
        s,
        "Exactly one of the two `print` calls runs — never both, never neither.",
        y,
        "ok",
    )

    s = slide("Example 2: the sign of a number", kicker="Worked example")
    y = code(
        s,
        [
            "# Read a number from the user",
            'num = int(input("Enter a number: "))',
            "",
            "# Determine if positive, negative, or zero",
            "if num > 0:",
            '    print("positive")',
            "elif num < 0:",
            '    print("negative")',
            "else:",
            '    print("zero")',
        ],
        w=Inches(6.6),
    )
    bullets(
        s,
        [
            "Three possible answers, so *three* branches.",
            "`num > 0` is tested first; only if it fails does Python try `num < 0`.",
            "`else` needs no condition — by then, the number can only be zero.",
            "Adding a fourth outcome means adding one more `elif`, nothing else.",
        ],
        x=M + Inches(6.9),
        w=BODY_W - Inches(6.9),
        y=BODY_TOP + Inches(0.1),
        size=16,
    )
    note(
        s,
        "Three outcomes, so three branches. `else` catches everything the "
        "earlier tests did not.",
        y + Inches(0.05),
        "tip",
    )

    s = slide("Example 3: classifying a triangle", kicker="Worked example")
    y = code(
        s,
        [
            "# Read three sides of a triangle",
            'side1 = float(input("Enter first side: "))',
            'side2 = float(input("Enter second side: "))',
            'side3 = float(input("Enter third side: "))',
            "",
            "# Classify the triangle",
            "if side1 == side2 == side3:",
            '    print("equilateral")',
            "elif side1 == side2 or side1 == side3 or side2 == side3:",
            '    print("isosceles")',
            "else:",
            '    print("scalene")',
        ],
        size=14,
    )
    bullets(
        s,
        [
            "`side1 == side2 == side3` is a chained comparison: all three equal.",
            "Equilateral is checked *first* — it is also isosceles by the second test.",
        ],
        y=y + Inches(0.05),
        size=17,
    )

    s = slide("Order of checks decides the answer", kicker="Order")
    y = code(
        s,
        [
            "# WRONG: every magnitude above 2.0 stops at the first branch",
            "if magnitude >= 2.0:",
            '    print("Very Minor")',
            "elif magnitude >= 8.0:",
            '    print("Great")        # never reached',
            "",
            "# RIGHT: start from the narrowest / highest range",
            "if magnitude >= 10.0:",
            '    print("Meteoric")',
            "elif magnitude >= 8.0:",
            '    print("Great")',
            "elif magnitude >= 2.0:",
            '    print("Very Minor")',
        ],
        size=14,
    )
    note(
        s,
        "Python takes the *first* branch that is `True` and skips all the "
        "rest. Sort your ranges before you write them.",
        y,
        "warn",
    )

    s = slide("elif vs. a stack of separate ifs", kicker="Order")
    y = code(
        s,
        [
            "# Two independent questions - both can print",
            "if n % 3 == 0:",
            '    print("Fizz")',
            "if n % 5 == 0:",
            '    print("Buzz")',
            "",
            "# One question with several answers - exactly one prints",
            "if n % 3 == 0 and n % 5 == 0:",
            '    print("FizzBuzz")',
            "elif n % 3 == 0:",
            '    print("Fizz")',
            "elif n % 5 == 0:",
            '    print("Buzz")',
            "else:",
            "    print(n)",
        ],
        size=13,
    )
    note(
        s,
        "Ask yourself: do the cases *exclude* each other? If yes, use "
        "`elif`. If no, use separate `if` statements.",
        y,
        "ok",
    )

    s = slide("Nested if statements", kicker="Nesting")
    y = code(
        s,
        [
            'num = float(input("Enter a number: "))',
            "",
            "if num > 0:",
            '    adjective = " "',
            "    if num >= 1000000:",
            '        adjective = " really big "',
            "    elif num >= 1000:",
            '        adjective = " big "',
            '    result = "That\'s a" + adjective + "positive number"',
            "elif num < 0:",
            '    result = "That\'s a negative number"',
            "else:",
            '    result = "That\'s zero"',
            "",
            "print(result)",
        ],
        size=13,
    )
    note(
        s,
        "The body of any branch may contain another `if`. Each nesting "
        "level adds 4 more spaces — past three levels, rethink the "
        "logic.",
        y,
        "tip",
    )

    s = slide("Part 4 in one slide", kicker="Summary")
    table(
        s,
        ["Construct", "What it does"],
        [
            ("`if`", "Runs a block only when a condition is true"),
            ("`if` / `else`", "Runs one of two blocks"),
            ("`if` / `elif` / `else`", "Runs one of several blocks, first match wins"),
            ("Nested `if`", "An `if` inside the body of another `if`"),
            ("Indentation", "Decides which lines belong to which branch"),
            ("Order of branches", "Changes the result whenever ranges overlap"),
        ],
        col_w=[4, 9],
        size=16,
    )
    note(
        s,
        "These four constructs let a program respond differently to every "
        "input — which is what the 39 exercises of this topic are "
        "about.",
        Inches(4.9),
        "ok",
    )


# ================================================ PART 5: CONDITIONS ON TEXT ==
def part5():
    section(
        "Conditions on Text",
        "Letters, words and the numbers hiding behind them",
        [
            "Comparing strings",
            "in and len",
            "Unpacking",
            "Case-insensitive input",
            "ord and chr",
            "ASCII",
        ],
    )

    s = slide("String operations you will need", kicker="Strings")
    table(
        s,
        ["Operation", "Example", "Purpose"],
        [
            (
                "`==`",
                '`"cat" == "dog"` -> `False`',
                "Compare strings for exact equality",
            ),
            (
                "`in`",
                '`"e" in "hello"` -> `True`',
                "Is this character or substring inside?",
            ),
            (
                "Unpacking",
                '`x, y, z = "cat"` -> `"c"`, `"a"`, `"t"`',
                "Split a short string into single characters",
            ),
            ("`len()`", '`len("hello")` -> `5`', "How many characters the string has"),
            ("`ord()`", '`ord("A")` -> `65`', "The ASCII number of a character"),
            ("`chr()`", '`chr(65)` -> `"A"`', "The character for an ASCII number"),
        ],
        col_w=[3, 6, 6],
        size=15,
    )
    note(
        s,
        '`in` is the shortest way to test a vowel: `if letter in "aeiou":`',
        Inches(5.1),
        "tip",
    )

    s = slide("String methods that answer questions", kicker="Strings")
    table(
        s,
        ["Method", "Example", "Purpose"],
        [
            (
                "`.lower()`",
                '`"Hello".lower()` -> `"hello"`',
                "Normalise before comparing",
            ),
            ("`.upper()`", '`"hello".upper()` -> `"HELLO"`', "The same, the other way"),
            (
                "`.title()`",
                '`"hello".title()` -> `"Hello"`',
                "Capitalise the first letter",
            ),
            (
                "`.isdigit()`",
                '`"5".isdigit()` -> `True`',
                "Is every character a digit?",
            ),
            (
                "`.isalpha()`",
                '`"a".isalpha()` -> `True`',
                "Is every character a letter?",
            ),
            ("`.isalnum()`", '`"a5".isalnum()` -> `True`', "Letters and digits only"),
            ("`.isspace()`", '`" ".isspace()` -> `True`', "Space, tab or newline"),
        ],
        col_w=[3, 6, 6],
        size=15,
    )
    note(
        s,
        "A method is called *on* a value: the dot belongs to the string, "
        "and the brackets are never optional.",
        Inches(5.45),
        "warn",
    )

    s = slide("The pattern for every text exercise", kicker="Strings")
    y = code(
        s,
        [
            "letter = input().lower()      # normalise once, right after reading",
            "",
            'if letter == "y":',
            '    print("sometimes vowel, sometimes consonant")',
            'elif letter in "aeiou":',
            '    print("vowel")',
            "else:",
            '    print("consonant")',
        ],
    )
    y = out(s, ["A", "vowel"], y=y)
    note(
        s,
        "Normalise the input *once*, then compare against lowercase literals only.",
        y,
        "ok",
    )

    s = slide("Splitting a short string into its characters", kicker="Strings")
    y = code(
        s,
        [
            '# A chess square such as "d5" is a letter and a digit',
            "position = input().lower()",
            'column, row = position        # unpacking: "d5" -> "d", "5"',
            "",
            'column_number = ord(column) - ord("a") + 1   # "a" -> 1 ... "h" -> 8',
            "row_number = int(row)",
            "",
            "if (column_number + row_number) % 2 == 0:",
            '    print("black")',
            "else:",
            '    print("white")',
        ],
        size=14,
    )
    note(
        s,
        "Unpacking needs the number of variables to match the number of "
        "characters exactly — otherwise you get a `ValueError`.",
        y,
        "warn",
    )

    s = slide("Behind every character is a number", kicker="ASCII")
    bullets(
        s,
        [
            "A computer stores text as numbers; ASCII is the oldest of those tables.",
            "`ord(c)` gives the number of a character, `chr(n)` gives the "
            "character back.",
            "Digits, capitals and lowercase letters each sit in one *continuous* "
            "block.",
            'That is why `"apple" < "banana"` works: Python compares the numbers.',
        ],
        size=18,
    )
    table(
        s,
        ["Range", "Characters", "Handy fact"],
        [
            ("48 – 57", "`0` – `9`", '`ord("7") - ord("0")` -> `7`'),
            ("65 – 90", "`A` – `Z`", '`ord("C") - ord("A")` -> `2`'),
            ("97 – 122", "`a` – `z`", "Lowercase is uppercase plus 32"),
            ("32", "space", '`ord(" ")` -> `32`'),
        ],
        y=Inches(4.3),
        col_w=[3, 4, 6],
        size=15,
    )
    note(
        s,
        "The full table is in `conditionals/_docs/ascii_table.md`.",
        Inches(6.35),
        "tip",
    )


# ======================================================= PART 6: MATCH ==
def part6():
    section(
        "The match Statement",
        "Python 3.10+ — clean branching on discrete values",
        [
            "Basic syntax",
            "Several values per case",
            "Wildcard",
            "Guards",
            "match vs if-elif",
        ],
    )

    s = slide("Basic match syntax", kicker="match")
    y = code(
        s,
        [
            "match expression:",
            "    case pattern1:",
            "        # code for pattern1",
            "    case pattern2:",
            "        # code for pattern2",
            "    case _:",
            "        # default case (optional)",
        ],
        size=17,
    )
    bullets(
        s,
        [
            "`match` compares *one* value against a list of patterns.",
            "The first matching `case` runs; the rest are skipped — just like `elif`.",
            "`case _:` is the wildcard: it matches anything the cases above did not.",
            "Requires Python *3.10 or newer* — this workbook uses 3.13.",
        ],
        y=y + Inches(0.1),
        size=18,
    )

    s = slide("The same task, two ways", kicker="match")
    code(
        s,
        [
            "month = input().lower()",
            "",
            "match month:",
            '    case "february":',
            '        print("28 or 29")',
            '    case "january" | "march" | "may" | "july" \\',
            '         | "august" | "october" | "december":',
            '        print("31")',
            '    case "april" | "june" | "september" | "november":',
            '        print("30")',
            "    case _:",
            '        print("Invalid month")',
        ],
        w=Inches(6.3),
        size=13,
    )
    code(
        s,
        [
            "month = input().lower()",
            "",
            'if month == "february":',
            '    print("28 or 29")',
            'elif month == "january" or month == "march" \\',
            '        or month == "may" or month == "july" \\',
            '        or month == "august" or month == "october" \\',
            '        or month == "december":',
            '    print("31")',
            'elif month == "april" or month == "june" \\',
            '        or month == "september" or month == "november":',
            '    print("30")',
            "else:",
            '    print("Invalid month")',
        ],
        x=M + Inches(6.6),
        w=BODY_W - Inches(6.6),
        size=13,
    )
    note(
        s,
        "Same result. The `|` in a `case` reads as *or* and keeps the "
        "branch on one line.",
        Inches(5.95),
        "tip",
    )

    s = slide("Matching plain values", kicker="match")
    y = code(
        s,
        [
            "# Read the denomination",
            "denomination = int(input())",
            "",
            "# Determine the individual on the banknote",
            "match denomination:",
            "    case 1:",
            '        print("Mirzo Tursunzoda")',
            "    case 20:",
            '        print("Abuali ibni Sino")',
            "    case 100:",
            '        print("Ismoili Somoni")',
            "    case 500:",
            '        print("Abuabdullo Rudaki")',
            "    case _:",
            '        print("Invalid denomination")',
        ],
        w=Inches(6.6),
        size=14,
    )
    bullets(
        s,
        [
            "A lookup table with a fixed set of keys is what `match` is best at.",
            "Numbers and strings both work as patterns.",
            "Without `case _`, an unknown value simply prints nothing — "
            "rarely what the task wants.",
        ],
        x=M + Inches(6.9),
        w=BODY_W - Inches(6.9),
        y=BODY_TOP + Inches(0.2),
        size=16,
    )

    s = slide("Guards: a case with an extra condition", kicker="match")
    y = code(
        s,
        [
            "# Read temperature",
            "temperature = int(input())",
            "",
            "# Classify temperature with guard conditions",
            "match temperature:",
            "    case temp if temp < 0:",
            '        print("freezing")',
            "    case temp if 0 <= temp < 20:",
            '        print("cold")',
            "    case temp if 20 <= temp < 30:",
            '        print("warm")',
            "    case temp if temp >= 30:",
            '        print("hot")',
        ],
        size=14,
    )
    note(
        s,
        "Guards make ranges possible, but at that point an `if` / `elif` "
        "chain is usually shorter and clearer.",
        y,
        "warn",
    )

    s = slide("match or if-elif?", kicker="match")
    table(
        s,
        ["Use", "When", "Example from the workbook"],
        [
            (
                "`match`",
                "One variable against several fixed values",
                "`faces_money`, `month_days`, `month_from_int`",
            ),
            (
                "`if` / `elif`",
                "Ranges, or conditions on several variables",
                "`richter_scale`, `valid_triangle`, `bmi_categories`",
            ),
            (
                "Either",
                "Small sets where both read well",
                "`vowel_consonant`, `student_year_level`",
            ),
        ],
        col_w=[3, 6, 6],
        size=16,
    )
    bullets(
        s,
        [
            "`match` is cleaner for *value-based* branching.",
            "`if` / `elif` is better for *condition-based* logic.",
            "Neither is more advanced than the other — pick the one that reads better.",
        ],
        y=Inches(4.3),
        size=18,
    )


# ================================================================ CLOSING ==
def closing():
    _section["name"] = ""

    s = slide("Your first five exercises", kicker="Homework", footer=False)
    rows = [
        ("1", "can_vote", "One `if` / `else` and a single comparison."),
        ("2", "even_odd", "The `%` operator inside a condition."),
        ("3", "number_sign", "Your first `elif` — three outcomes."),
        ("4", "vowel_consonant", "Text input, `.lower()` and `in`."),
        ("5", "two_equal", "Combining comparisons with `or`."),
    ]
    y = BODY_TOP
    for num, name, desc in rows:
        _box(
            s,
            M,
            y,
            Inches(0.55),
            Inches(0.55),
            fill=YELLOW,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.15,
        )
        tfn = _tf(s, M, y, Inches(0.55), Inches(0.55), anchor=MSO_ANCHOR.MIDDLE)
        _para(tfn, num, 18, NAVY, bold=True, first=True, align=PP_ALIGN.CENTER)
        tft = _tf(s, M + Inches(0.8), y - Inches(0.02), Inches(10.5), Inches(0.62))
        _para(tft, name, 19, NAVY, bold=True, font=MONO, first=True, space_after=1)
        p = tft.add_paragraph()
        _rich(p, _markup(desc), 15, GRAY)
        y += Inches(0.78)
    y = note(
        s,
        "Run one exercise with "
        "`uv run pytest conditionals/can_vote/`, or the whole topic "
        "with `uv run pytest conditionals/`.",
        y + Inches(0.1),
        "ok",
    )
    note(
        s,
        "All 39 exercises, in order: `conditionals/_docs/_table_of_contents.md`",
        y,
        "tip",
    )

    s = slide("Cheat sheet — everything on one page", kicker="Reference", footer=False)
    table(
        s,
        ["Task", "Code"],
        [
            (
                "Compare two values",
                "`a == b`   `a != b`   `a < b`   `a <= b`   `a > b`   `a >= b`",
            ),
            ("Run a block only if true", "`if age >= 18:`"),
            ("Two branches", "`if ...:` / `else:`"),
            ("Many branches", "`if ...:` / `elif ...:` / `else:`"),
            ("Both must hold", "`if a > 0 and b > 0:`"),
            ("At least one holds", "`if a == b or a == c:`"),
            ("Flip a condition", "`if not is_open:`"),
            ("A range", "`if 18.5 <= bmi < 25.0:`"),
            ("All three equal", "`if a == b == c:`"),
            ("Ignore letter case", "`month = input().lower()`"),
            ("Is a character in a set", '`if letter in "aeiou":`'),
            ("Character to number and back", '`ord("A")`   `chr(65)`'),
            ("Branch on fixed values", "`match value:` / `case 1:` / `case _:`"),
        ],
        col_w=[5, 9],
        size=14,
    )

    s = slide("Where to read more", kicker="Reference", footer=False)
    table(
        s,
        ["Document", "What is in it"],
        [
            ("`comparison_operators.md`", "Part 2 — the six operators"),
            (
                "`boolean_operators.md`",
                "Part 3 — truth tables, precedence, short-circuit",
            ),
            ("`if_elif_else.md`", "Part 4 — the chapter text and worked examples"),
            ("`string_utils.md`", "Part 5 — string operations and methods"),
            ("`ascii_table.md`", "Part 5 — the full ASCII table"),
            ("`match.md`", "Part 6 — the match statement"),
            ("`_table_of_contents.md`", "Reading order + all 39 exercises"),
        ],
        col_w=[6, 8],
        size=16,
    )
    note(s, "All of them live in `conditionals/_docs/`.", Inches(5.3), "tip")

    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, 0, W, Inches(0.22), fill=YELLOW)
    tf = _tf(s, Inches(1.3), Inches(2.2), Inches(10.7), Inches(3.2))
    _para(
        tf,
        "Now your programs can choose",
        42,
        WHITE,
        bold=True,
        first=True,
        space_after=18,
    )
    _para(
        tf,
        "Start with can_vote, and read the task statement twice before "
        "you write a line.",
        22,
        RGBColor(0xC3, 0xD3, 0xE1),
        space_after=26,
    )
    _para(
        tf,
        "When a test fails, compare the expected and the actual output "
        "character by character — the answer is almost always "
        "there.",
        17,
        RGBColor(0x8F, 0xB6, 0xD9),
    )
    tf2 = _tf(s, Inches(1.3), Inches(5.6), Inches(10.7), Inches(0.8))
    _para(tf2, "Good luck  ·  see you soon", 20, YELLOW, bold=True, first=True)


# ==================================================================== main ==
title_slide()
roadmap_slide()
part1()
part2()
part3()
part4()
part5()
part6()
closing()

prs.save(OUT)
print(f"{len(prs.slides)} slides -> {OUT}")
